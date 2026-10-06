"""
Adds to the final file the molecules from final_database.csv that are not
already in it (matched by CAS number). One row is added per new molecule:
  - the BlueBird / method_1 row if its retention time (RT) is not 0
  - otherwise, the first row from another column with RT != 0
    (any method, following FALLBACK_ORDER)
Molecules with no usable RT anywhere are skipped and listed.
The original file is left untouched: the result is written to OUTPUT.
 
Dependencies: pip install pandas openpyxl
"""
 
from pathlib import Path
 
import pandas as pd
from openpyxl import load_workbook
 
# ======================= CONFIGURATION =======================
# Script lives in <project>/src/dev, data in <project>/data
PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = PROJECT_ROOT / "data"
 
DATABASE = DATA_DIR / "final_database.csv"
FINAL_FILE = DATA_DIR / "Fichier final (RT+sol).xlsx"
OUTPUT = DATA_DIR / "Fichier final (RT+sol)_MAJ.xlsx"
 
# Preferred row
PREFERRED = ("BlueBird", "method_1")
 
# Fallback order when the preferred row has RT = 0 (first match with RT != 0 wins)
FALLBACK_ORDER = [
    ("PFP", "method_1"), ("PolarTec", "method_1"), ("RP18", "method_1"), ("Sphinx", "method_1"),
    ("PFP", "method_2"), ("PolarTec", "method_2"), ("RP18", "method_2"), ("Sphinx", "method_2"),
    ("PFP", "method_3"), ("PolarTec", "method_3"), ("RP18", "method_3"),
]
 
# Column mapping: database column -> final file column
MAPPING = {
    "column": "Column",
    "method": "Method",
    "Sample Name": "Sample Name",
    "CAS #": "CAS ",            # note: trailing space in the final file
    "RT": "RT",
    "Signal": "Signal",
    "Matrix ID Name": "Matrix ID Name",
}
CAS_COL_DB = "CAS #"
CAS_COL_FINAL = "CAS "
MATRIX_PREFIX = "matrix ID "    # "PFP 1" -> "matrix ID PFP 1"
# =============================================================
 
 
def norm_cas(x):
    return None if pd.isna(x) else str(x).strip()
 
 
def pick_row(rows):
    """Return (row, used_fallback) for one molecule, or (None, False) if no usable RT."""
    def find(column, method):
        m = rows[(rows["_col"] == column) & (rows["_met"] == method) & (rows["_rt"] != 0)
                 & rows["_rt"].notna()]
        return None if m.empty else m.iloc[0]
 
    row = find(*PREFERRED)
    if row is not None:
        return row, False
    for column, method in FALLBACK_ORDER:
        row = find(column, method)
        if row is not None:
            return row, True
    return None, False
 
 
def main():
    db = pd.read_csv(DATABASE, sep=None, engine="python")
    final = pd.read_excel(FINAL_FILE)
 
    db["_cas"] = db[CAS_COL_DB].map(norm_cas)
    db["_col"] = db["column"].astype(str).str.strip()
    db["_met"] = db["method"].astype(str).str.strip()
    db["_rt"] = pd.to_numeric(db["RT"], errors="coerce")
    db = db.dropna(subset=["_cas"])
 
    # 1) Molecules not yet in the final file (by CAS)
    existing = set(final[CAS_COL_FINAL].map(norm_cas).dropna())
    new_cas = [c for c in db["_cas"].unique() if c not in existing]
    print(f"Molecules in database      : {db['_cas'].nunique()}")
    print(f"Already in final file      : {db['_cas'].nunique() - len(new_cas)}")
 
    # 2) Pick one row per new molecule
    selected, fallbacks, no_rt = [], [], []
    for cas in new_cas:
        rows = db[db["_cas"] == cas]
        row, used_fallback = pick_row(rows)
        if row is None:
            no_rt.append((rows["Sample Name"].iloc[0], cas))
            continue
        selected.append(row)
        if used_fallback:
            fallbacks.append((row["Sample Name"], cas, row["_col"], row["_met"], row["_rt"]))
 
    # 3) Append to the Excel file (original formatting preserved)
    wb = load_workbook(FINAL_FILE)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    for row in selected:
        values = {col_final: row[col_db] for col_db, col_final in MAPPING.items()}
        values["Matrix ID Name"] = MATRIX_PREFIX + str(row["Matrix ID Name"]).strip()
        values[CAS_COL_FINAL] = row["_cas"]
        out = []
        for h in headers:
            v = values.get(h)
            out.append(None if pd.isna(v) else (v.item() if hasattr(v, "item") else v))
        ws.append(out)
    wb.save(OUTPUT)
 
    # Summary
    print("\n===== SUMMARY =====")
    print(f"New molecules added        : {len(selected)}")
    print(f"  - {PREFERRED[0]} / {PREFERRED[1]}      : {len(selected) - len(fallbacks)}")
    print(f"  - other column (RT = 0 on {PREFERRED[0]}): {len(fallbacks)}")
    for name, cas, col, met, rt in fallbacks:
        print(f"      {name} ({cas}) -> {col} / {met}, RT = {rt}")
    if no_rt:
        print(f"Skipped (RT = 0 everywhere): {len(no_rt)}")
        for name, cas in no_rt:
            print(f"      {name} ({cas})")
    print("Note: solubility cells (MeOH, ACN, DMSO, DCM, CHCl3) are empty for added rows.")
    print(f"Saved to: {OUTPUT}")
 
 
if __name__ == "__main__":
    main()
 
